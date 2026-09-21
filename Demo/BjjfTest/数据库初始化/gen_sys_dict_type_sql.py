# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_dict_type 插入语句（不含自增主键 dict_id）

用法:
  python gen_sys_dict_type_sql.py
  python gen_sys_dict_type_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
import argparse
import re
import os

DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "sql_export-be148300-2438-44b5-8d1c-e61f1caf31d9-260720140845.xlsx",
)
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Desktop", "sys_dict_type_insert.sql")

COLS = [
    "dict_name",
    "dict_type",
    "status",
    "create_by",
    "create_time",
    "update_by",
    "update_time",
    "remark",
]

DATETIME_COLS = {"create_time", "update_time"}
NULLABLE_STR_COLS = {"remark"}


def sql_str(v):
    if v is None:
        return "NULL"
    s = str(v)
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_datetime(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def format_value(col, v):
    if col in DATETIME_COLS:
        return sql_datetime(v)
    if col in NULLABLE_STR_COLS:
        if v is None or v == "":
            return "NULL"
        return sql_str(v)
    if v is None:
        defaults = {
            "dict_name": "",
            "dict_type": "",
            "status": "0",
            "create_by": "",
            "update_by": "",
        }
        return sql_str(defaults.get(col, ""))
    return sql_str(v)


def generate(xlsx_path: str, out_path: str) -> int:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in COLS if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    col_list = ", ".join(f"`{c}`" for c in COLS)
    lines = [
        "-- sys_dict_type data export",
        "-- without dict_id (auto increment handled by MySQL)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    seen_types = set()
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        dict_type = row[idx["dict_type"]]
        if dict_type is None or str(dict_type).strip() == "":
            continue
        dict_type_s = str(dict_type).strip()
        # UNIQUE KEY dict_type：跳过重复
        if dict_type_s in seen_types:
            continue
        seen_types.add(dict_type_s)

        values = [format_value(c, row[idx[c]]) for c in COLS]
        # dict_type 用去空白后的值
        values[COLS.index("dict_type")] = sql_str(dict_type_s)
        sql = f"INSERT INTO `sys_dict_type` ({col_list}) VALUES (" + ", ".join(values) + ");"
        lines.append(sql)
        count += 1

    wb.close()

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    parser = argparse.ArgumentParser(description="Export sys_dict_type INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        raise SystemExit(f"xlsx not found: {args.xlsx}")

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
