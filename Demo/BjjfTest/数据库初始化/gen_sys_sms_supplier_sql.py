# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_sms_supplier 插入语句（不含自增主键 id）

执行链路:
  1. 读桌面（或 --xlsx）Excel，第一行为表头
  2. 按列名定位字段，逐行转成 INSERT
  3. 跳过空行 / 无 supplier_number 的行
  4. 写出 SET NAMES + INSERT 到 .sql

用法:
  python gen_sys_sms_supplier_sql.py
  python gen_sys_sms_supplier_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import argparse
import re
import os

DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "sql_export-60bcdbd6-5f8d-4f0a-9c63-3ed4d2c77211-260722093816.xlsx",
)
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Desktop", "sys_sms_supplier_insert.sql")

# 不含 id：自增主键由库生成
COLS = [
    "supplier_name",
    "supplier_number",
    "mobile_code_cost",
    "unicom_code_cost",
    "telecom_code_cost",
    "mobile_notice_cost",
    "unicom_notice_cost",
    "telecom_notice_cost",
    "mobile_market_cost",
    "unicom_market_cost",
    "telecom_market_cost",
    "contain_domain",
    "contain_sign",
    "create_time",
    "update_time",
    "is_delete",
    "mobile_sign",
    "unicom_sign",
    "telecom_sign",
]

DECIMAL_COLS = {
    "mobile_code_cost",
    "unicom_code_cost",
    "telecom_code_cost",
    "mobile_notice_cost",
    "unicom_notice_cost",
    "telecom_notice_cost",
    "mobile_market_cost",
    "unicom_market_cost",
    "telecom_market_cost",
}
TINYINT_COLS = {"contain_domain", "contain_sign", "is_delete"}
DATETIME_COLS = {"create_time", "update_time"}
NULLABLE_STR_COLS = {"supplier_name", "mobile_sign", "unicom_sign", "telecom_sign"}


def sql_str(v):
    """字符串 → SQL 字面量；None → NULL。"""
    if v is None:
        return "NULL"
    s = str(v)
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_decimal(v):
    """decimal(11,4)；空 → NULL。"""
    if v is None or v == "":
        return "NULL"
    try:
        d = Decimal(str(v).strip())
    except (InvalidOperation, ValueError):
        return "NULL"
    return format(d, "f")


def sql_tinyint(v):
    """tinyint；空按 0。"""
    if v is None or v == "":
        return "0"
    return str(int(float(str(v))))


def sql_datetime(v, nullable=True):
    """
    日期时间 → DATETIME 字面量。
    create_time 不可空：空则 CURRENT_TIMESTAMP；update_time 可空。
    """
    if v is None or v == "":
        return "NULL" if nullable else "CURRENT_TIMESTAMP"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def format_value(col, v):
    if col in DECIMAL_COLS:
        return sql_decimal(v)
    if col in TINYINT_COLS:
        return sql_tinyint(v)
    if col == "create_time":
        return sql_datetime(v, nullable=False)
    if col == "update_time":
        return sql_datetime(v, nullable=True)
    if col in NULLABLE_STR_COLS:
        if v is None or str(v).strip() == "":
            return "NULL"
        return sql_str(v)
    # supplier_number：业务编号，空串兜底
    if v is None:
        return sql_str("")
    return sql_str(v)


def generate(xlsx_path: str, out_path: str) -> int:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in COLS if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    col_list = ", ".join(f"`{c}`" for c in COLS)
    lines = [
        "-- sys_sms_supplier data export",
        "-- without id (auto increment handled by MySQL)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        # supplier_number 作为有效行判据
        number = row[idx["supplier_number"]]
        if number is None or str(number).strip() == "":
            continue

        values = [format_value(c, row[idx[c]]) for c in COLS]
        sql = f"INSERT INTO `sys_sms_supplier` ({col_list}) VALUES (" + ", ".join(values) + ");"
        lines.append(sql)
        count += 1

    wb.close()

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    parser = argparse.ArgumentParser(description="Export sys_sms_supplier INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        raise SystemExit(f"xlsx not found: {args.xlsx}")

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
