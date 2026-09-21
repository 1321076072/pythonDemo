# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_dict_data 插入语句（不含自增主键 dict_code）

用法:
  python gen_sys_dict_data_sql.py
  python gen_sys_dict_data_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
import argparse
import re
import os

DEFAULT_XLSX = r"C:\Users\useradmin\Desktop\sql_export-3bf8660f-e366-447e-89a2-9ed124411470-260720141242.xlsx"
DEFAULT_OUT = r"C:\Users\useradmin\Desktop\sys_dict_data_insert.sql"

COLS = [
    "dict_sort",
    "dict_label",
    "dict_value",
    "dict_type",
    "css_class",
    "list_class",
    "is_default",
    "status",
    "create_by",
    "create_time",
    "update_by",
    "update_time",
    "remark",
]

DATETIME_COLS = {"create_time", "update_time"}
NULLABLE_STR_COLS = {"css_class", "list_class", "remark"}
INT_COLS = {"dict_sort"}


def sql_str(v):
    if v is None:
        return "NULL"
    s = str(v)
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_int(v, default=0):
    if v is None or v == "":
        return str(default)
    return str(int(float(str(v))))


def sql_datetime(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def format_value(col, v):
    if col in INT_COLS:
        return sql_int(v, 0)
    if col in DATETIME_COLS:
        return sql_datetime(v)
    if col in NULLABLE_STR_COLS:
        if v is None or v == "":
            return "NULL"
        return sql_str(v)
    if v is None:
        defaults = {
            "dict_label": "",
            "dict_value": "",
            "dict_type": "",
            "is_default": "N",
            "status": "0",
            "create_by": "",
            "update_by": "",
        }
        return sql_str(defaults.get(col, ""))
    return sql_str(v)


def generate(xlsx_path: str, out_path: str) -> int:
    if not os.path.isfile(xlsx_path):
        raise SystemExit(f"xlsx not found: {xlsx_path}")

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in COLS if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    col_list = ", ".join(f"`{c}`" for c in COLS)
    lines = [
        "-- sys_dict_data data export",
        "-- without dict_code (auto increment handled by MySQL)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        dict_type = row[idx["dict_type"]]
        dict_label = row[idx["dict_label"]]
        if dict_type is None or str(dict_type).strip() == "":
            continue
        if dict_label is None or str(dict_label).strip() == "":
            continue

        values = [format_value(c, row[idx[c]]) for c in COLS]
        sql = f"INSERT INTO `sys_dict_data` ({col_list}) VALUES (" + ", ".join(values) + ");"
        lines.append(sql)
        count += 1

    wb.close()

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    parser = argparse.ArgumentParser(description="Export sys_dict_data INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
