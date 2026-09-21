# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_config 的 INSERT SQL（不含 config_id，交给库自增）。

执行链路（现象 → 本质）:
  1. 读桌面（或 --xlsx）上的 Excel，第一行为表头
  2. 按列名定位字段，逐行转成 MySQL INSERT
  3. 跳过空行 / 无 config_key 的行（无效配置）
  4. 写出 SET NAMES + 多条 INSERT 到 .sql（或 --out）

用法:
  python gen_sys_config_sql.py
  python gen_sys_config_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
import argparse
import re
import os

# ---------- 默认路径：桌面导出的 xlsx → 桌面生成的 sql ----------
DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "sql_export-85196293-5a2d-401f-9398-8bf4548caba0-260720095443.xlsx",
)
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Desktop", "sys_config_insert.sql")


def sql_str(v):
    """
    字符串 → SQL 字面量。
    None → NULL；其余转 str，转义反斜杠与单引号后包上引号。
    """
    if v is None:
        return "NULL"
    s = str(v)
    # MySQL 字符串：\ → \\，' → ''
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_tinyint(v):
    """
    布尔/删除标记类字段 → tinyint。
    空值按 0；支持 "1.0" 这类 Excel 浮点数字符串。
    """
    if v is None or v == "":
        return "0"
    return str(int(float(str(v))))


def sql_datetime(v, nullable=True):
    """
    日期时间 → SQL DATETIME 字面量。
    - 空：nullable 为 True 写 NULL，否则写 CURRENT_TIMESTAMP（create_time 用）
    - datetime/date 对象：格式化为 'YYYY-MM-DD HH:MM:SS'
    - 纯日期字符串：自动补 ' 00:00:00'
    - 其它字符串：走 sql_str 转义
    """
    if v is None or v == "":
        return "NULL" if nullable else "CURRENT_TIMESTAMP"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    # Excel 有时给出 ISO：2024-01-01T12:00:00 → 空格分隔
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def generate(xlsx_path: str, out_path: str) -> int:
    """
    核心：xlsx → INSERT 列表 → 写文件。
    返回成功生成的 INSERT 条数。
    """
    # data_only=True：读单元格计算后的值，不读公式本身
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    # 第 0 行当表头；建 name → 列下标，后面按名取值，不依赖列顺序
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    # 目标表字段（刻意不含 config_id：自增主键由库生成，避免环境冲突）
    cols = [
        "config_name",
        "config_key",
        "config_value",
        "config_type",
        "create_by",
        "update_by",
        "remark",
        "create_time",
        "update_time",
        "is_delete",
    ]
    missing = [c for c in cols if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    # SQL 文件头：注释 + 字符集，保证中文配置不被乱码
    lines = [
        "-- sys_config data export",
        "-- without config_id (auto increment handled by MySQL)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    for row in rows[1:]:
        # 整行空 → 跳过（导出尾部常见空行）
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        # config_key 是业务主键语义；没有 key 的行不是有效配置
        key = row[idx["config_key"]]
        if key is None or str(key).strip() == "":
            continue

        name = row[idx["config_name"]]
        value = row[idx["config_value"]]
        cfg_type = row[idx["config_type"]]
        create_by = row[idx["create_by"]]
        update_by = row[idx["update_by"]]
        remark = row[idx["remark"]]
        create_time = row[idx["create_time"]]
        update_time = row[idx["update_time"]]
        is_delete = row[idx["is_delete"]]

        # 各字段空值策略不同：
        # - name/create_by/update_by：空串，避免 NOT NULL 列插 NULL
        # - value/remark：允许 NULL
        # - config_type：缺省 'N'（若依惯例：Y=内置 N=自定义）
        # - create_time：不可空，空则 CURRENT_TIMESTAMP
        # - update_time：可空
        # - is_delete：空按 0
        values = [
            sql_str("" if name is None else name),
            sql_str(key),
            "NULL" if value is None else sql_str(value),
            sql_str("N" if cfg_type is None or cfg_type == "" else cfg_type),
            sql_str("" if create_by is None else create_by),
            sql_str("" if update_by is None else update_by),
            "NULL" if remark is None else sql_str(remark),
            sql_datetime(create_time, nullable=False),
            sql_datetime(update_time, nullable=True),
            sql_tinyint(is_delete),
        ]

        sql = (
            "INSERT INTO `sys_config` ("
            "`config_name`, `config_key`, `config_value`, `config_type`, "
            "`create_by`, `update_by`, `remark`, `create_time`, `update_time`, `is_delete`"
            ") VALUES (" + ", ".join(values) + ");"
        )
        lines.append(sql)
        count += 1

    wb.close()

    # 统一 LF，避免 Windows CRLF 在 Linux 上执行时出问题
    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    """入口：解析参数 → 校验文件存在 → generate → 打印条数。"""
    parser = argparse.ArgumentParser(description="Export sys_config INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        raise SystemExit(f"xlsx not found: {args.xlsx}")

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
