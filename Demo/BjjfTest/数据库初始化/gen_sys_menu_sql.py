# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_menu 插入语句（包含自增主键 menu_id）

用法:
  python gen_sys_menu_sql.py
  python gen_sys_menu_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
import argparse
import re
import os

DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "sql_export-73c69274-4d7f-4dbd-9828-b13a18c0bbba-260720103440.xlsx",
)
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Desktop", "sys_menu_insert.sql")

COLS = [
    "menu_id",
    "menu_name",
    "parent_id",
    "order_num",
    "path",
    "component",
    "is_frame",
    "is_cache",
    "menu_type",
    "visible",
    "status",
    "perms",
    "icon",
    "create_by",
    "create_time",
    "update_by",
    "update_time",
    "remark",
]

INT_COLS = {"menu_id", "parent_id", "order_num", "is_frame", "is_cache"}
NULLABLE_STR_COLS = {"component", "perms"}
DATETIME_COLS = {"create_time", "update_time"}


def sql_str(v):
    if v is None:
        return "NULL"
    s = str(v)
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_int(v, default=None):
    if v is None or v == "":
        if default is None:
            return "NULL"
        return str(default)
    return str(int(float(str(v))))


def sql_datetime(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def format_value(col, v):
    if col in INT_COLS:
        defaults = {
            "menu_id": None,
            "parent_id": 0,
            "order_num": 0,
            "is_frame": 1,
            "is_cache": 0,
        }
        return sql_int(v, defaults.get(col))
    if col in DATETIME_COLS:
        return sql_datetime(v)
    if col in NULLABLE_STR_COLS:
        if v is None or v == "":
            return "NULL"
        return sql_str(v)
    # char / varchar with empty default
    if v is None:
        if col in ("icon",):
            return "'#'"
        if col in ("visible", "status"):
            return "'0'"
        return "''"
    return sql_str(v)


def generate(xlsx_path: str, out_path: str) -> int:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    missing = [c for c in COLS if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    col_list = ", ".join(f"`{c}`" for c in COLS)
    lines = [
        "-- sys_menu data export",
        "-- includes menu_id (preserve primary keys)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    max_menu_id = 0
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        menu_id = row[idx["menu_id"]]
        menu_name = row[idx["menu_name"]]
        if menu_id is None or menu_name is None or str(menu_name).strip() == "":
            continue

        values = [format_value(c, row[idx[c]]) for c in COLS]
        sql = f"INSERT INTO `sys_menu` ({col_list}) VALUES (" + ", ".join(values) + ");"
        lines.append(sql)
        count += 1
        try:
            mid = int(float(str(menu_id)))
            if mid > max_menu_id:
                max_menu_id = mid
        except Exception:
            pass

    wb.close()

    # 导入后校正自增起点，避免后续新增主键冲突
    if max_menu_id > 0:
        lines.append("")
        lines.append(f"ALTER TABLE `sys_menu` AUTO_INCREMENT = {max_menu_id + 1};")

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    parser = argparse.ArgumentParser(description="Export sys_menu INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        raise SystemExit(f"xlsx not found: {args.xlsx}")

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
