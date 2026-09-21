# -*- coding: utf-8 -*-
"""
从 xlsx 导出 sys_job 插入语句（不含自增主键 job_id）

用法:
  python gen_sys_job_sql.py
  python gen_sys_job_sql.py --xlsx "路径.xlsx" --out "路径.sql"
"""
from openpyxl import load_workbook
from datetime import datetime, date
import argparse
import re
import os

DEFAULT_XLSX = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "sql_export-83d645c1-4fd0-43aa-a88d-5125d008c1cf-260720112804.xlsx",
)
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Desktop", "sys_job_insert.sql")

# 不含 job_id
COLS = [
    "job_name",
    "job_group",
    "invoke_target",
    "cron_expression",
    "misfire_policy",
    "concurrent",
    "status",
    "create_by",
    "create_time",
    "update_by",
    "update_time",
    "remark",
]

DATETIME_COLS = {"create_time", "update_time"}


def sql_str(v):
    if v is None:
        return "NULL"
    s = str(v)
    s = s.replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def sql_datetime(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip().replace("T", " ")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        s = s + " 00:00:00"
    return sql_str(s)


def format_value(col, v):
    if col in DATETIME_COLS:
        return sql_datetime(v)
    if v is None:
        defaults = {
            "job_group": "DEFAULT",
            "cron_expression": "",
            "misfire_policy": "3",
            "concurrent": "1",
            "status": "0",
            "create_by": "",
            "update_by": "",
            "remark": "",
        }
        if col in defaults:
            return sql_str(defaults[col])
        return "NULL"
    # trim trailing spaces on names commonly present in export
    if col == "job_name":
        return sql_str(str(v).rstrip())
    return sql_str(v)


def generate(xlsx_path: str, out_path: str) -> int:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    needed = COLS + ["job_id"]  # job_id 在表头中存在即可，导出时不写
    missing = [c for c in COLS if c not in idx]
    if missing:
        raise SystemExit(f"missing columns: {missing}, header={header}")

    col_list = ", ".join(f"`{c}`" for c in COLS)
    lines = [
        "-- sys_job data export",
        "-- without job_id (auto increment handled by MySQL)",
        "-- source xlsx rows converted to INSERT",
        "",
        "SET NAMES utf8mb4;",
        "",
    ]

    count = 0
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        job_name = row[idx["job_name"]]
        invoke_target = row[idx["invoke_target"]]
        if job_name is None or str(job_name).strip() == "":
            continue
        if invoke_target is None or str(invoke_target).strip() == "":
            continue

        values = [format_value(c, row[idx[c]]) for c in COLS]
        sql = f"INSERT INTO `sys_job` ({col_list}) VALUES (" + ", ".join(values) + ");"
        lines.append(sql)
        count += 1

    wb.close()

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
        f.write("\n")

    return count


def main():
    parser = argparse.ArgumentParser(description="Export sys_job INSERT SQL from xlsx")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="xlsx file path")
    parser.add_argument("--out", default=DEFAULT_OUT, help="output .sql path")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        raise SystemExit(f"xlsx not found: {args.xlsx}")

    count = generate(args.xlsx, args.out)
    print(f"written: {args.out}")
    print(f"insert_count: {count}")


if __name__ == "__main__":
    main()
